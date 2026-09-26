#!/usr/bin/env python3
"""Enrich ranked + directory CSVs, build reference tables, docs helpers, xlsx, JSON."""
from __future__ import annotations

import csv
import json
import re
from collections import Counter, OrderedDict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
RAW = ROOT / "raw_sources"
DOCS = ROOT / "docs"
TODAY = "2026-09-25"

RANKED_NEW_COLS = [
    "Website", "LinkedIn_URL", "Twitter_X_URL", "Founders", "CEO_or_Lead",
    "HQ_City", "HQ_Country", "Operating_Countries", "Legal_Entity_Name",
    "Company_Status", "Exit_Type", "Acquirer", "Exit_Year", "Exit_Source_URL", "Business_Model", "Target_Customers", "Core_Products",
    "Tech_Stack_Hints", "Employee_Range", "Latest_Round_Type",
    "Latest_Round_Amount_USD", "Latest_Round_Date", "Notable_Investors",
    "Licence_Types_Held", "Key_Competitors", "Moat_Notes",
    "Replication_Capital_Intensity", "Replication_Time_Estimate", "Risk_Flags",
    "Builder_Notes", "Data_Confidence",
]

DIR_NEW_COLS = [
    "Website", "Regulator_Source", "Licence_or_Category", "Licence_Status",
    "HQ_City", "HQ_Country", "Business_Model", "Heuristic",
    "Builder_Notes", "Data_Confidence",
]

# Category heuristics for directory Ease / licences / opportunity
CATEGORY_HEURISTICS = {
    "CBN PSSP": {
        "ease": 2, "licences": "CBN PSSP; PCI-DSS; NDPR; CAC.",
        "bm": "B2B payments", "opp": "Sell reconciliation, dispute automation, or merchant plugins to PSSPs.",
        "cat": "PSSP",
    },
    "CBN PTSP": {
        "ease": 2, "licences": "CBN PTSP; device certification; NDPR; CAC.",
        "bm": "B2B POS/terminals", "opp": "Sell terminal fleet analytics or merchant CRM to PTSPs.",
        "cat": "PTSP",
    },
    "CBN PTSA": {
        "ease": 2, "licences": "CBN PTSA; NDPR; CAC.",
        "bm": "B2B terminal aggregation", "opp": "Sell monitoring and settlement tooling to terminal aggregators.",
        "cat": "PTSA",
    },
    "CBN Super-Agent": {
        "ease": 2, "licences": "CBN Super-Agent; agent network compliance; NDPR; CAC.",
        "bm": "B2B2C agency banking", "opp": "Sell agent performance and AML monitoring to Super-Agents.",
        "cat": "Super-Agent",
    },
    "CBN Mobile Money Operator (MMO)": {
        "ease": 1, "licences": "CBN MMO; NDIC where applicable; NDPR; CAC.",
        "bm": "B2C mobile money", "opp": "Sell wallet analytics and agent tooling; do not attempt unlicensed e-money.",
        "cat": "MMO",
    },
    "CBN Switching & Processing": {
        "ease": 1, "licences": "CBN Switching and Processing; PCI-DSS; NDPR.",
        "bm": "B2B switching", "opp": "Build adjacent merchant UX; national switch replication is unrealistic for small teams.",
        "cat": "Switching",
    },
    "CBN Card/Payment Scheme": {
        "ease": 1, "licences": "CBN Card/Payment Scheme; scheme rules; PCI-DSS.",
        "bm": "B2B card scheme", "opp": "Sell issuing/acquiring analytics to scheme participants.",
        "cat": "Card Scheme",
    },
    "CBN Payments Service Holding": {
        "ease": 1, "licences": "CBN Payments Service Holding company rules; subsidiary licences.",
        "bm": "Holding / payments group", "opp": "Group-level compliance and reporting tooling.",
        "cat": "Payments Holding",
    },
    "CBN International Money Transfer Operator (IMTO)": {
        "ease": 2, "licences": "CBN IMTO; AML/CFT; partner bank; NDPR.",
        "bm": "B2B/B2C remittance", "opp": "Sell payout orchestration and compliance case management to IMTOs.",
        "cat": "IMTO",
    },
    "FCCPC-registered Digital Money Lender (DML)": {
        "ease": 2, "licences": "FCCPC Digital Money Lender registration; credit bureau; NDPR; consumer protection.",
        "bm": "B2C digital lending", "opp": "Sell ethical collections, underwriting features, or model monitoring to DMLs.",
        "cat": "Digital Money Lender",
    },
}

SECTOR_DEFAULTS = {
    "Fintech": {"ease": 2, "licences": "CAC; NDPR; payments/lending licences as applicable.", "bm": "Fintech", "opp": "Offer automation/ML to fintech ops (recon, KYC, collections)."},
    "SaaS": {"ease": 4, "licences": "CAC; NDPR.", "bm": "B2B SaaS", "opp": "High fit for productized consultancy builds and integrations."},
    "Marketplace": {"ease": 3, "licences": "CAC; NDPR; consumer protection; payments via licensed PSPs.", "bm": "Marketplace", "opp": "Sell trust/fraud, catalog, or logistics tooling to marketplaces."},
    "Logistics": {"ease": 3, "licences": "CAC; transport permits; insurance; NDPR.", "bm": "Logistics", "opp": "Sell routing, dispatch, or telematics analytics."},
    "Healthtech": {"ease": 3, "licences": "CAC; health data rules; NDPR; facility/NAFDAC as applicable.", "bm": "Healthtech", "opp": "Sell clinical NLP, claims, or interoperability connectors."},
    "Edtech": {"ease": 4, "licences": "CAC; NDPR (minors); education standards if school-integrated.", "bm": "Edtech", "opp": "Sell assessment analytics or content tooling."},
    "Agritech": {"ease": 3, "licences": "CAC; NDPR; lending partners / food safety as applicable.", "bm": "Agritech", "opp": "Sell satellite verification, credit models, or marketplace ops tools."},
    "Cleantech": {"ease": 2, "licences": "CAC; NERC-adjacent; energy/credit partners.", "bm": "Cleantech / energy", "opp": "Sell monitoring, PAYG, or credit tooling to energy operators."},
    "AI": {"ease": 4, "licences": "CAC; NDPR.", "bm": "AI / software", "opp": "Partner or compete on vertical AI for Nigerian ops workflows."},
    "Insurtech": {"ease": 2, "licences": "NAICOM; NDPR; CAC.", "bm": "Insurtech", "opp": "Sell claims NLP and fraud detection to HMOs/insurers."},
    "Other Tech": {"ease": 3, "licences": "CAC; NDPR.", "bm": "Other tech", "opp": "Evaluate product vs services; sell automation where software-native."},
}

JUNK_DIRECTORY_EXACT = {
    "London Iron House London, SE1 1UN United Kingdom",
    "Maclay Murray & Spens LLP 1 George Square, Glasgow, G2 1AL, United Kingdom",
    "No 1, Kimihurura, Gasabo, Umujyi Wa Kigali, Rwanda",
    "Portal 2, The Beacon, Westgate, Newscastle Upon Tyne, UK. NE4 9PQ",
    "Intergrated Revenue Collection and Management Solution",  # typo dup of Integrated...
}

DIRECTORY_RENAMES = {
    "18. Xpress Payments Solution": "Xpress Payments Solution",
}


def read_csv(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, rows: list[dict], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in fieldnames})


def parse_hq(hq: str) -> tuple[str, str]:
    hq = (hq or "").strip()
    if not hq or hq.lower() in {"unknown", "nan"}:
        return "", ""
    # crude split
    if "Nigeria" in hq and "/" not in hq and "," not in hq:
        return hq if hq != "Nigeria" else "", "Nigeria"
    if "Lagos" in hq:
        country = "Nigeria" if "Nigeria" in hq or "Lagos" in hq else ""
        if "London" in hq or "UK" in hq:
            return "Lagos / London", "Nigeria / UK"
        return "Lagos", "Nigeria"
    if "Abuja" in hq:
        return "Abuja", "Nigeria"
    if "Ibadan" in hq:
        return "Ibadan", "Nigeria"
    if "Kaduna" in hq:
        return "Kaduna", "Nigeria"
    return hq, ""


def merge_sources(*parts: str) -> str:
    urls = []
    seen = set()
    for p in parts:
        if not p:
            continue
        for u in re.split(r"[;\n|]+", str(p)):
            u = u.strip()
            if u and u not in seen and u.lower() not in {"nan", "none"}:
                seen.add(u)
                urls.append(u)
    return "; ".join(urls)


def enrich_ranked() -> tuple[list[dict], list[str]]:
    rows = read_csv(DATA / "ranked_disclosed.csv")
    enrich = json.loads((DATA / "ranked_enrichment.json").read_text())
    # also pull YC websites
    yc = {x["name"].lower(): x for x in json.loads((RAW / "yc_nigeria_names.json").read_text())}
    out = []
    base_cols = list(rows[0].keys())
    # desired column order
    cols = []
    for c in [
        "Rank", "Company", "Sector", "Product_Summary", "Website", "LinkedIn_URL", "Twitter_X_URL",
        "Founders", "CEO_or_Lead", "HQ_or_Primary_Market", "HQ_City", "HQ_Country", "Operating_Countries",
        "Legal_Entity_Name", "Founded_Year", "Company_Status", "Exit_Type", "Acquirer", "Exit_Year", "Exit_Source_URL", "Business_Model", "Target_Customers",
        "Core_Products", "Tech_Stack_Hints", "Employee_Range",
        "Total_Disclosed_Funding_USD", "Latest_Round_Type", "Latest_Round_Amount_USD", "Latest_Round_Date",
        "Notable_Investors", "Valuation_USD", "Valuation_Status",
        "Licence_Types_Held", "Licences_Regs_Needed", "Key_Competitors", "Moat_Notes",
        "Ease_to_Replicate", "Replicate_Notes", "Replication_Capital_Intensity", "Replication_Time_Estimate",
        "Risk_Flags", "Builder_Notes", "Data_Confidence", "Source_URLs", "Last_Verified_Date",
    ]:
        if c not in cols:
            cols.append(c)

    for r in rows:
        name = r["Company"]
        e = enrich.get(name, {})
        city, country = parse_hq(r.get("HQ_or_Primary_Market", ""))
        # YC website fallback
        website = e.get("Website", "")
        if not website:
            for yk, yv in yc.items():
                if yk in name.lower() or name.lower() in yk:
                    website = yv.get("website", "") or website
                    break
        row = dict(r)
        for c in RANKED_NEW_COLS:
            row[c] = e.get(c, row.get(c, ""))
        if website:
            row["Website"] = website
        if not row.get("HQ_City"):
            row["HQ_City"] = e.get("HQ_City") or city
        if not row.get("HQ_Country"):
            row["HQ_Country"] = e.get("HQ_Country") or country
        if not row.get("Core_Products") and r.get("Product_Summary"):
            row["Core_Products"] = ""  # keep blank unless researched; Product_Summary remains
        if not row.get("Company_Status"):
            row["Company_Status"] = "Unknown"
        if not row.get("Data_Confidence"):
            row["Data_Confidence"] = "Low"
        # improve replicate notes from enrichment moat if empty improvement desired
        if e.get("Moat_Notes") and (not row.get("Replicate_Notes") or len(str(row.get("Replicate_Notes", ""))) < 40):
            pass  # keep existing Replicate_Notes; Moat_Notes is separate
        if e.get("Replication_Capital_Intensity") and not row.get("Replication_Capital_Intensity"):
            row["Replication_Capital_Intensity"] = e["Replication_Capital_Intensity"]
        # merge sources
        row["Source_URLs"] = merge_sources(r.get("Source_URLs", ""), e.get("Source_URLs", ""))
        row["Last_Verified_Date"] = TODAY
        # fill Licences_Regs_Needed improvement if enrichment has Licence_Types_Held and existing thin
        out.append({k: ("" if row.get(k) is None else str(row.get(k))) for k in cols})

    write_csv(DATA / "ranked_disclosed.csv", out, cols)
    return out, cols


def detect_regulator(product: str, sources: str) -> tuple[str, str, str]:
    text = f"{product} {sources}"
    for key, h in CATEGORY_HEURISTICS.items():
        short = key.replace("CBN ", "").replace("FCCPC-registered ", "")
        if key in text or short in text:
            # map
            if "FCCPC" in key or "Digital Money" in text:
                return "FCCPC", h["cat"], "Listed"
            if "IMTO" in key or "International Money Transfer" in text:
                return "CBN", "IMTO", "Listed"
            if "MMO" in key or "Mobile Money" in text:
                return "CBN", "MMO", "Listed"
            if "PSSP" in key or "Payment Solution Service" in text or "Payment Service Provider" in text:
                return "CBN", "PSSP", "Listed"
            if "PTSP" in key or "Payment Terminal Service Provider" in text:
                return "CBN", "PTSP", "Listed"
            if "Super-Agent" in key or "Super Agent" in text:
                return "CBN", "Super-Agent", "Listed"
            if "Switching" in key or "Switching" in text:
                return "CBN", "Switching", "Listed"
            if "Card" in key:
                return "CBN", "Card Scheme", "Listed"
            return "CBN" if "CBN" in key else "FCCPC", h["cat"], "Listed"
    if "Y Combinator" in text or "ycombinator" in text.lower() or "YC" in sources:
        return "YC", "YC Startup", "Listed"
    if "startupmapafrica" in sources.lower() or "StartupMapAfrica" in text:
        return "Press", "", "Unknown"
    if "startuplist" in sources.lower():
        return "Press", "", "Unknown"
    return "Other", "", "Unknown"


def enrich_directory(ranked_names: set[str]) -> tuple[list[dict], list[str]]:
    rows = read_csv(DATA / "directory.csv")
    yc = {x["name"].lower(): x for x in json.loads((RAW / "yc_nigeria_names.json").read_text())}
    curated = {x["name"].lower(): x for x in json.loads((DATA / "curated_directory.json").read_text())}

    out = []
    for r in rows:
        name = (r.get("Company") or "").strip()
        if name in JUNK_DIRECTORY_EXACT:
            continue
        if name in DIRECTORY_RENAMES:
            name = DIRECTORY_RENAMES[name]
            r["Company"] = name
            # fix product blurb leading number
            ps = r.get("Product_Summary", "")
            r["Product_Summary"] = re.sub(r"^\d+\.\s*", "", ps)

        # skip if exact ranked duplicate
        # keep directory exclusive; ranked already separate

        product = r.get("Product_Summary", "")
        sources = r.get("Source_URLs", "")
        sector = r.get("Sector", "Other Tech")
        reg, lic_cat, lic_status = detect_regulator(product, sources)

        # improve thin SMA blurbs
        if "Listed on StartupMapAfrica" in product:
            product = f"{name}: Nigeria-focused {sector.lower()} company listed on StartupMapAfrica sector directories. Funding undisclosed in free sources used."
            r["Product_Summary"] = product

        heur = CATEGORY_HEURISTICS.get(
            next((k for k in CATEGORY_HEURISTICS if k.split("(")[0].strip() in f"{product} {lic_cat}" or CATEGORY_HEURISTICS[k]["cat"] == lic_cat), ""),
            None,
        )
        if not heur:
            # try by detected cat
            heur = next((v for v in CATEGORY_HEURISTICS.values() if v["cat"] == lic_cat), None)
        if not heur:
            heur = SECTOR_DEFAULTS.get(sector, SECTOR_DEFAULTS["Other Tech"])
            # normalize sector default shape
            if "cat" not in heur:
                heur = {
                    "ease": heur["ease"],
                    "licences": heur["licences"],
                    "bm": heur["bm"],
                    "opp": heur["opp"],
                    "cat": lic_cat or sector,
                }

        city, country = parse_hq(r.get("HQ_or_Primary_Market", ""))
        website = ""
        y = yc.get(name.lower())
        if y:
            website = y.get("website", "")
            reg = reg if reg != "Other" else "YC"
        c = curated.get(name.lower())
        if c and not website:
            # curated may not have website
            pass

        # For pure regulator rows, leave Website blank (per brief)
        if reg in {"CBN", "FCCPC"} and not website:
            website = ""

        ease = r.get("Ease_to_Replicate") or heur["ease"]
        licences = r.get("Licences_Regs_Needed") or heur["licences"]
        # Prefer category-specific licence text when regulator known
        if reg in {"CBN", "FCCPC"} and heur.get("licences"):
            licences = heur["licences"]

        conf = "Med" if reg in {"CBN", "FCCPC", "YC"} else "Low"
        source = merge_sources(sources)
        if reg == "CBN" and "cbn.gov.ng" not in source.lower():
            # keep existing; many already cite CBN
            pass

        row = dict(r)
        row["Company"] = name
        row["Product_Summary"] = product
        row["Website"] = website
        row["Regulator_Source"] = reg
        row["Licence_or_Category"] = lic_cat or heur.get("cat", "")
        row["Licence_Status"] = lic_status if lic_cat else "Unknown"
        row["HQ_City"] = city or ("Lagos" if "Lagos" in str(r.get("HQ_or_Primary_Market", "")) else "")
        row["HQ_Country"] = country or ("Nigeria" if "Nigeria" in str(r.get("HQ_or_Primary_Market", "")) or reg in {"CBN", "FCCPC"} else "")
        row["Business_Model"] = heur.get("bm", "")
        row["Ease_to_Replicate"] = ease
        row["Licences_Regs_Needed"] = licences
        row["Heuristic"] = "Yes"
        row["Builder_Notes"] = heur.get("opp", "")
        row["Data_Confidence"] = conf
        row["Source_URLs"] = source
        row["Last_Verified_Date"] = TODAY
        row["Replicate_Notes"] = r.get("Replicate_Notes") or (
            f"Directory heuristic for {row['Licence_or_Category'] or sector}: see docs/methodology.md. Not a researched deep-dive."
        )
        out.append(row)

    # dedupe by normalized name keeping first
    seen = set()
    deduped = []
    for r in out:
        key = re.sub(r"[^a-z0-9]+", "", r["Company"].lower())
        if key in seen:
            continue
        seen.add(key)
        deduped.append(r)

    cols = [
        "Company", "Sector", "Product_Summary", "Website", "Regulator_Source", "Licence_or_Category",
        "Licence_Status", "HQ_or_Primary_Market", "HQ_City", "HQ_Country", "Founded_Year",
        "Business_Model", "Total_Disclosed_Funding_USD", "Valuation_USD", "Valuation_Status",
        "Ease_to_Replicate", "Replicate_Notes", "Licences_Regs_Needed", "Heuristic",
        "Builder_Notes", "Data_Confidence", "Source_URLs", "Notes", "Last_Verified_Date",
    ]
    write_csv(DATA / "directory.csv", deduped, cols)
    return deduped, cols


def build_funding_and_investors(ranked: list[dict]) -> tuple[list[dict], list[dict]]:
    rounds = []
    # Only include rounds where we have type or amount from enrichment (real sourced)
    for r in ranked:
        rt = (r.get("Latest_Round_Type") or "").strip()
        amt = (r.get("Latest_Round_Amount_USD") or "").strip()
        dt = (r.get("Latest_Round_Date") or "").strip()
        if not rt and not amt:
            continue
        if rt.lower() in {"ipo / public markets"} and not amt:
            rounds.append({
                "Company": r["Company"],
                "Round_Type": rt,
                "Amount_USD": "",
                "Date": dt,
                "Investors": r.get("Notable_Investors", ""),
                "Source_URLs": r.get("Source_URLs", ""),
                "Notes": "Public company / IPO path; not a private round amount.",
            })
            continue
        if amt or (rt and dt):
            rounds.append({
                "Company": r["Company"],
                "Round_Type": rt,
                "Amount_USD": amt,
                "Date": dt,
                "Investors": r.get("Notable_Investors", ""),
                "Source_URLs": r.get("Source_URLs", ""),
                "Notes": "Latest disclosed round fields from public sources; may not be complete history.",
            })

    # Known extra historical rounds (only well-sourced)
    EXTRA = [
        {"Company": "Moniepoint (TeamApt)", "Round_Type": "Series C (first close)", "Amount_USD": "110000000", "Date": "2024-10",
         "Investors": "Development Partners International; Google Africa Investment Fund; Verod; Lightrock",
         "Source_URLs": "https://techcrunch.com/2024/10/29/google-dpi-backs-moniepoint-in-110m-round/",
         "Notes": "First close of Series C; later completed >$200M (Oct 2025)."},
        {"Company": "Flutterwave", "Round_Type": "Series D", "Amount_USD": "250000000", "Date": "2022-01",
         "Investors": "Tiger Global; Avenir Growth; others (press)",
         "Source_URLs": "https://www.cbinsights.com/company/flutterwave/financials",
         "Notes": "Widely reported Series D."},
        {"Company": "OPay", "Round_Type": "Series C", "Amount_USD": "400000000", "Date": "2021-08",
         "Investors": "SoftBank Vision Fund 2; others (press)",
         "Source_URLs": "https://ln247.news/top-10-most-funded-nigerian-startups-as-of-august-2023/",
         "Notes": "Large Series C widely reported."},
        {"Company": "Yellow Card", "Round_Type": "Series C", "Amount_USD": "33000000", "Date": "2024-10",
         "Investors": "Blockchain Capital; Polychain; Third Prime; Castle Island; Block; Galaxy; Winklevoss Capital; Hutt Capital",
         "Source_URLs": "https://yellowcard.io/blog/yellow-card-closes-us-33m-series-c-funding-round-to-drive-global-expansion-and-strategic-initiatives",
         "Notes": ""},
        {"Company": "LemFi", "Round_Type": "Series B", "Amount_USD": "53000000", "Date": "2025-01",
         "Investors": "Highland Europe; Left Lane Capital; Palm Drive Capital; Y Combinator; Endeavor Catalyst",
         "Source_URLs": "https://techcrunch.com/2025/01/13/lemfi-moves-remittances-further-into-asia-and-europe-with-53m-in-new-funding/",
         "Notes": ""},
        {"Company": "Helium Health", "Round_Type": "Growth", "Amount_USD": "30000000", "Date": "2023-06",
         "Investors": "AXA IM; Anne Wojcicki (press)",
         "Source_URLs": "https://techcrunch.com/2023/06/05/helium-health-gets-30m-backed-by-axa-im-and-23andmes-anne-wojcicki/",
         "Notes": ""},
        {"Company": "Reliance Health", "Round_Type": "Series B", "Amount_USD": "40000000", "Date": "2022-02",
         "Investors": "General Atlantic",
         "Source_URLs": "https://www.businesswire.com/news/home/20220207005155/en/Reliance-Health-Raises-%2440M-in-Series-B-Led-by-General-Atlantic",
         "Notes": ""},
        {"Company": "Nomba", "Round_Type": "Growth", "Amount_USD": "30000000", "Date": "2023-05",
         "Investors": "Base10 Partners; Shopify",
         "Source_URLs": "https://techcrunch.com/2023/05/02/african-payment-service-provider-nomba-raises-30m-backed-by-base10-partners-and-shopify/",
         "Notes": ""},
        {"Company": "Paystack", "Round_Type": "Acquisition", "Amount_USD": "", "Date": "2020-10",
         "Investors": "Stripe",
         "Source_URLs": "https://paystack.com",
         "Notes": "Acquired by Stripe; ~$200M valuation often cited in press for the deal."},
        {"Company": "Mono", "Round_Type": "Acquisition", "Amount_USD": "", "Date": "2026-01",
         "Investors": "Flutterwave",
         "Source_URLs": "https://mono.co/blog/flutterwave-mono-acquisition",
         "Notes": "All-stock acquisition; press cited deal value up to ~$40M."},
        {"Company": "Gokada", "Round_Type": "Series A", "Amount_USD": "5300000", "Date": "2019-05",
         "Investors": "",
         "Source_URLs": "https://techcrunch.com/2019/05/24/nigerias-gokada-raises-5-3m-round-for-its-motorcycle-ride-hail-biz/",
         "Notes": ""},
        {"Company": "NowNow", "Round_Type": "Seed", "Amount_USD": "13000000", "Date": "2022-09",
         "Investors": "",
         "Source_URLs": "https://techcabal.com/2022/09/07/nownow-raises-13-million-in-seed-funding-to-expand-services-across-africa/",
         "Notes": ""},
        {"Company": "Mecho Autotech", "Round_Type": "Pre-Series A", "Amount_USD": "2400000", "Date": "2023-09",
         "Investors": "",
         "Source_URLs": "https://techcabal.com/2023/09/14/nigerian-startup-mecho-autotech-raises-2-4-million-in-pre-series-a-round/",
         "Notes": ""},
        {"Company": "Vendease", "Round_Type": "Growth", "Amount_USD": "30000000", "Date": "2022-09",
         "Investors": "",
         "Source_URLs": "https://techcabal.com/2022/09/26/vendease-raises-30-million-to-offer-procurement-services-across-africa/",
         "Notes": ""},
        {"Company": "Kuda", "Round_Type": "Series B", "Amount_USD": "55000000", "Date": "2021-08",
         "Investors": "Target Global; Valar Ventures (press)",
         "Source_URLs": "https://kuda.com",
         "Notes": ""},
    ]
    # dedupe by company+date+amount
    seen = set()
    all_rounds = []
    for rr in rounds + EXTRA:
        key = (rr["Company"], rr.get("Date", ""), str(rr.get("Amount_USD", "")), rr.get("Round_Type", ""))
        if key in seen:
            continue
        seen.add(key)
        all_rounds.append(rr)

    inv_counter = Counter()
    inv_sources = {}
    inv_companies = {}
    for rr in all_rounds:
        for part in re.split(r"[;]", rr.get("Investors") or ""):
            name = part.strip()
            name = re.sub(r"\s*\(.*?\)\s*", "", name).strip()
            if not name or name.lower() in {"others", "press", "public shareholders", "others (press)", "y combinator (historical)"}:
                continue
            if "press" in name.lower() and len(name) < 20:
                continue
            inv_counter[name] += 1
            inv_sources.setdefault(name, rr.get("Source_URLs", ""))
            inv_companies.setdefault(name, set()).add(rr["Company"])

    # also scan Notable_Investors on ranked
    for r in ranked:
        for part in re.split(r"[;]", r.get("Notable_Investors") or ""):
            name = part.strip()
            name = re.sub(r"\s*\(.*?\)\s*", "", name).strip()
            if not name or len(name) < 2:
                continue
            if name.lower() in {"others", "press", "undisclosed in this compile for full list; see press roundups"}:
                continue
            inv_counter[name] += 0  # ensure present
            inv_sources.setdefault(name, r.get("Source_URLs", ""))
            inv_companies.setdefault(name, set()).add(r["Company"])

    def inv_type(n: str) -> str:
        nl = n.lower()
        if any(x in nl for x in ["y combinator", "ycombinator"]):
            return "Accelerator"
        if any(x in nl for x in ["ifc", "bii", "proparco", "swedfund", "fmo", "google africa"]):
            return "DFI / CVC"
        if any(x in nl for x in ["visa", "shopify", "stripe", "softbank", "tencent", "transsion"]):
            return "CVC / Strategic"
        if any(x in nl for x in ["capital", "ventures", "partners", "fund", "investment", "qed", "tiger", "sequoia", "accel", "general atlantic", "leapfrog", "lightrock", "verod", "novastar"]):
            return "VC/PE"
        if "angel" in nl:
            return "Angel"
        return "Other / Unknown"

    investors = []
    for name, _ in sorted(inv_counter.items(), key=lambda x: (-x[1], x[0].lower())):
        investors.append({
            "Investor": name,
            "Type": inv_type(name),
            "Focus": "Africa / Nigeria tech (as appears in this dataset)",
            "Example_Companies": "; ".join(sorted(inv_companies.get(name, []))[:8]),
            "Source_URLs": inv_sources.get(name, ""),
        })

    write_csv(DATA / "funding_rounds.csv", all_rounds,
              ["Company", "Round_Type", "Amount_USD", "Date", "Investors", "Source_URLs", "Notes"])
    write_csv(DATA / "investors.csv", investors,
              ["Investor", "Type", "Focus", "Example_Companies", "Source_URLs"])
    return all_rounds, investors


def write_ranked_json(ranked: list[dict]) -> None:
    arr = []
    for r in ranked:
        arr.append({
            "rank": int(r["Rank"]) if str(r["Rank"]).isdigit() else r["Rank"],
            "company": r["Company"],
            "sector": r["Sector"],
            "product_summary": r.get("Product_Summary", ""),
            "website": r.get("Website", ""),
            "linkedin_url": r.get("LinkedIn_URL", ""),
            "twitter_x_url": r.get("Twitter_X_URL", ""),
            "founders": [x.strip() for x in (r.get("Founders") or "").split(";") if x.strip()],
            "ceo_or_lead": r.get("CEO_or_Lead", ""),
            "hq": {
                "display": r.get("HQ_or_Primary_Market", ""),
                "city": r.get("HQ_City", ""),
                "country": r.get("HQ_Country", ""),
            },
            "operating_countries": [x.strip() for x in (r.get("Operating_Countries") or "").split(";") if x.strip()],
            "legal_entity_name": r.get("Legal_Entity_Name", ""),
            "founded_year": r.get("Founded_Year", ""),
            "company_status": r.get("Company_Status", ""),
            "business_model": r.get("Business_Model", ""),
            "target_customers": r.get("Target_Customers", ""),
            "core_products": [x.strip() for x in (r.get("Core_Products") or "").split(";") if x.strip()],
            "tech_stack_hints": r.get("Tech_Stack_Hints", ""),
            "employee_range": r.get("Employee_Range", ""),
            "funding": {
                "total_disclosed_usd": r.get("Total_Disclosed_Funding_USD", ""),
                "latest_round_type": r.get("Latest_Round_Type", ""),
                "latest_round_amount_usd": r.get("Latest_Round_Amount_USD", ""),
                "latest_round_date": r.get("Latest_Round_Date", ""),
                "notable_investors": [x.strip() for x in (r.get("Notable_Investors") or "").split(";") if x.strip()],
            },
            "valuation": {
                "usd": r.get("Valuation_USD", ""),
                "status": r.get("Valuation_Status", ""),
            },
            "licences": {
                "types_held": r.get("Licence_Types_Held", ""),
                "regs_needed": r.get("Licences_Regs_Needed", ""),
            },
            "competition": {
                "key_competitors": r.get("Key_Competitors", ""),
                "moat_notes": r.get("Moat_Notes", ""),
            },
            "replication": {
                "ease_to_replicate": r.get("Ease_to_Replicate", ""),
                "replicate_notes": r.get("Replicate_Notes", ""),
                "capital_intensity": r.get("Replication_Capital_Intensity", ""),
                "time_estimate": r.get("Replication_Time_Estimate", ""),
            },
            "risk_flags": r.get("Risk_Flags", ""),
            "builder_notes": r.get("Builder_Notes", ""),
            "data_confidence": r.get("Data_Confidence", ""),
            "source_urls": [x.strip() for x in (r.get("Source_URLs") or "").split(";") if x.strip()],
            "last_verified_date": r.get("Last_Verified_Date", TODAY),
        })
    (DATA / "companies_enriched.json").write_text(json.dumps(arr, indent=2))


def style_header(ws):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def autosize(ws, max_width=48):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col[:80]:
            length = max(length, min(len(str(cell.value or "")), max_width))
        ws.column_dimensions[letter].width = max(12, length + 2)


def write_xlsx(ranked, ranked_cols, directory, dir_cols, rounds, investors):
    wb = Workbook()
    # Ranked
    ws = wb.active
    ws.title = "Ranked_Disclosed"
    ws.append(ranked_cols)
    for r in ranked:
        ws.append([r.get(c, "") for c in ranked_cols])
    style_header(ws)
    autosize(ws)

    ws2 = wb.create_sheet("Directory")
    ws2.append(dir_cols)
    for r in directory:
        ws2.append([r.get(c, "") for c in dir_cols])
    style_header(ws2)
    autosize(ws2)

    # No disclosed notable
    notable = read_csv(DATA / "no_disclosed_funding.csv")
    if notable:
        cols = list(notable[0].keys())
        ws3 = wb.create_sheet("No_Disclosed_Funding")
        ws3.append(cols)
        for r in notable:
            ws3.append([r.get(c, "") for c in cols])
        style_header(ws3)
        autosize(ws3)

    ws4 = wb.create_sheet("Funding_Rounds")
    fcols = ["Company", "Round_Type", "Amount_USD", "Date", "Investors", "Source_URLs", "Notes"]
    ws4.append(fcols)
    for r in rounds:
        ws4.append([r.get(c, "") for c in fcols])
    style_header(ws4)
    autosize(ws4)

    ws5 = wb.create_sheet("Investors")
    icols = ["Investor", "Type", "Focus", "Example_Companies", "Source_URLs"]
    ws5.append(icols)
    for r in investors:
        ws5.append([r.get(c, "") for c in icols])
    style_header(ws5)
    autosize(ws5)

    # Data dictionary summary sheet
    dd = [
        ("Rank", "Integer rank by Total_Disclosed_Funding_USD descending"),
        ("Company", "Common trading / brand name"),
        ("Total_Disclosed_Funding_USD", "Sum of publicly disclosed funding we could source; never invented"),
        ("Valuation_USD", "Public valuation figure or Undisclosed"),
        ("Data_Confidence", "High / Med / Low based on source quality"),
        ("Heuristic", "Directory only: Yes if Ease/licences from category defaults"),
        ("Licence_Types_Held", "Licences company is reported to hold (verify on regulator sites)"),
        ("Builder_Notes", "Optional research notes on adjacent tooling (not a sales pipeline)"),
    ]
    ws6 = wb.create_sheet("Data_Dictionary")
    ws6.append(["Column", "Definition"])
    for a, b in dd:
        ws6.append([a, b])
    style_header(ws6)
    autosize(ws6)
    ws6.append([])
    ws6.append(["See docs/data_dictionary.md for the full column list."])

    # Methodology brief
    ws7 = wb.create_sheet("Methodology")
    for line in [
        ["Topic", "Note"],
        ["Ranking", "Primary sort: Total_Disclosed_Funding_USD descending. Ties broken by existing order."],
        ["No invention", "Funding, valuations, employee counts, founders, licence status left blank/Unknown/Undisclosed when not public."],
        ["Debt vs equity", "Some press totals mix equity and debt. See Replicate_Notes / company Notes (e.g. Moove, VertoFX)."],
        ["FX", "Figures kept in USD as reported by sources. Naira rounds converted only when source stated USD."],
        ["Directory heuristics", "Ease_to_Replicate and default licences for CBN/FCCPC rows use category defaults in scripts/enrich_and_rebuild.py. Heuristic=Yes."],
        ["Licences", "Not legal advice. Always re-check CBN/FCCPC/SEC/NAICOM primary pages before commercial reliance."],
        ["Last verified", TODAY],
    ]:
        ws7.append(line)
    style_header(ws7)
    autosize(ws7)

    licences = read_csv(DATA / "licences_cheat_sheet.csv")
    if licences:
        ws8 = wb.create_sheet("Licences_Cheat_Sheet")
        cols = list(licences[0].keys())
        ws8.append(cols)
        for r in licences:
            ws8.append([r.get(c, "") for c in cols])
        style_header(ws8)
        autosize(ws8)

    out1 = DATA / "Nigeria_Fintech_Tech_Ranked.xlsx"
    out2 = ROOT / "Nigeria_Fintech_Tech_Ranked.xlsx"
    wb.save(out1)
    wb.save(out2)


def coverage_stats(ranked, directory):
    def filled(rows, col):
        return sum(1 for r in rows if str(r.get(col, "")).strip() and str(r.get(col, "")).strip().lower() not in {"nan", "undisclosed", "unknown"})

    n = len(ranked)
    stats = {
        "ranked_count": n,
        "directory_count": len(directory),
        "ranked_website": filled(ranked, "Website"),
        "ranked_founders": filled(ranked, "Founders"),
        "ranked_ceo": filled(ranked, "CEO_or_Lead"),
        "ranked_linkedin": filled(ranked, "LinkedIn_URL"),
        "ranked_latest_round": filled(ranked, "Latest_Round_Type"),
        "ranked_investors": filled(ranked, "Notable_Investors"),
        "ranked_status": filled(ranked, "Company_Status"),
        "ranked_builder_notes": filled(ranked, "Builder_Notes"),
        "ranked_confidence_high": sum(1 for r in ranked if r.get("Data_Confidence") == "High"),
        "dir_website": filled(directory, "Website"),
        "dir_regulator": filled(directory, "Regulator_Source"),
        "dir_licence_cat": filled(directory, "Licence_or_Category"),
        "dir_builder_notes": filled(directory, "Builder_Notes"),
        "last_verified": TODAY,
    }
    (DATA / "build_meta.json").write_text(json.dumps(stats, indent=2))
    return stats


def main():
    ranked, ranked_cols = enrich_ranked()
    ranked_names = {r["Company"] for r in ranked}
    directory, dir_cols = enrich_directory(ranked_names)
    rounds, investors = build_funding_and_investors(ranked)
    write_ranked_json(ranked)
    write_xlsx(ranked, ranked_cols, directory, dir_cols, rounds, investors)
    stats = coverage_stats(ranked, directory)
    print(json.dumps(stats, indent=2))
    print("funding_rounds", len(rounds), "investors", len(investors))
    print("directory after clean", len(directory))


if __name__ == "__main__":
    main()
