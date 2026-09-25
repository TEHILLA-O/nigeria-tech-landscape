#!/usr/bin/env python3
"""Expand Nigeria tech landscape CSVs/XLSX from licence lists + curated seeds."""
from __future__ import annotations
import csv, json, re
from collections import OrderedDict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA, RAW = ROOT / "data", ROOT / "raw_sources"
TODAY = "2026-09-25"

ALIASES = {
    "teamapt": "moniepoint", "teamapt limited": "moniepoint",
    "moniepoint teamapt": "moniepoint", "appzone": "zone",
    "zone payment network limited formerly appzone limited": "zone",
    "payhippo": "rivy", "paylater": "carbon", "paylater hub limited": "carbon",
    "opay digital services limited formerly paycom nigeria limited": "opay",
    "opay digital services limited": "opay", "pagatech limited": "paga",
    "paga remit limited formerly pagatech limited": "paga",
    "nownow digital systems limited formerly contec global infotech limited": "nownow",
    "nownow digital systems limited": "nownow",
    "kongapay technologies limited formerly zinternet limited": "kongapay",
    "mkudi limited": "nomba",
    "nomba financial services limited formerly cosmic intelligence lab limited": "nomba",
    "nomba financial services limited": "nomba",
    "flutterwave technology solutions limited": "flutterwave",
    "flutterwave tech payments limited formerly flutterwave technology solutions ltd": "flutterwave",
    "paystack payment limited": "paystack", "interswitch limited": "interswitch",
    "palmpay limited": "palmpay", "abeg technologies limited": "abeg",
    "identitypass": "prembly", "prembly formerly identitypass": "prembly",
    "helicarrier prev buycoins": "buycoins", "helicarrier": "buycoins",
    "rank formerly moni": "rank", "trade depot limited": "tradedepot",
    "metro africa xpress max": "max ng", "max ng": "max ng",
    "swwipe financial services limited formerly parkway projects limited": "parkway",
    "parkway projects limited": "parkway",
    "rightcard payment services ltd a lemfi brand company": "lemfi",
    "aladdin scheme limited": "aladdin digital bank",
    "montra technology solutions limited formerly artha fintech limited": "montra",
    "venture garden nigeria limited": "venture garden group",
    "transfercorp limited vfd group": "vfd group",
    "midddleman technologies": "middleman", "eja ice": "eja ice", "eja-ice": "eja ice",
    "10mg health": "10mg health", "10mg health 2": "10mg health",
    "touch and pay technologies limited": "touch pay",
    "blaaiz innvoations technology limited": "blaaiz", "vertofx ltd": "vertofx",
    "waza technologies inc": "waza", "raenest inc": "raenest", "afriex inc": "afriex",
    "fincra technologies limited": "fincra", "duplo limited": "duplo",
    "klasha technologies limited": "klasha", "okra technologies limited": "okra",
    "cellulant nigeria limited": "cellulant", "onepipe io services ltd": "onepipe",
    "eyowo integrated payments limited": "eyowo", "clane company nig ltd": "clane",
    "3line card management limited": "3line", "3line card management": "3line",
    "nigeria inter bank settlement system plc": "nibss",
    "accelerex networks limited": "global accelerex", "global accelerex limited": "global accelerex",
    "smile identity nigeria limited": "smile identity",
    "remita payment service limited": "remita",
    "unified payment services limited": "unified payments",
    "hydrogen payment services limited": "hydrogen",
    "coralpay technology nigeria limited": "coralpay",
    "kredete technology limited": "kredete", "payaza africa limited": "payaza",
    "go lemon": "golemon", "buy coins": "buycoins",
}

SKIP = {
    "visa international service association", "mastercard international",
    "american express international", "unionpay international",
    "mastercard transaction services mts", "western union", "moneygram",
    "remitly inc", "ria financials", "transferto mobile financial services limited thunes",
    "taptap send uk limited", "small world financial services group limited",
    "idt payment services inc", "paysend plc", "chime inc sendwave",
    "volopa financial services scotland limited", "network international",
    "nigerian postal service nipost", "airtel mobile commerce nigeria limited airtel",
    "page not found", "everything you need one membership",
    "sourcing deals in this market",
}

def norm_key(name: str) -> str:
    s = name.lower().strip()
    s = re.sub(r"\(.*?\)", " ", s)
    s = s.replace("&", " and ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    s = re.sub(r"\b(limited|ltd|plc|inc|llc|corp|corporation|nigeria|nig|technologies|technology|tech|solutions|services|service|financial|finance|payments|payment|digital|systems|system|company|co|group|africa|international|integrations|integrated)\b", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    if s in ALIASES: return ALIASES[s]
    raw = re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", name.lower())).strip()
    if raw in ALIASES: return ALIASES[raw]
    return s or name.lower().strip()

def brand_from_legal(name: str) -> str:
    s = re.sub(r"\s*\(.*?\)\s*", " ", name.strip())
    s = re.sub(r"\b(Limited|Ltd\.?|PLC|Inc\.?|LLC|Nigeria|Nig\.?)\b", "", s, flags=re.I)
    s = re.sub(r"\s+", " ", s).strip(" ,.-")
    if s.isupper() and len(s) > 3: s = s.title()
    return s or name.strip()

def should_skip(name: str) -> bool:
    k = name.lower().strip()
    k2 = re.sub(r"[^a-z0-9]+", " ", k).strip()
    if k in SKIP or k2 in SKIP: return True
    if "western union" in k or "moneygram" in k or "remitly" in k: return True
    if re.search(r"\b(street|road|avenue|suite|plaza|building|freeway|lane)\b", k) and not re.search(r"\b(limited|ltd|inc|plc)\b", k):
        return True
    if re.match(r"^\d+\.?\s*$", k) or len(k) < 3: return True
    return False

def product_for(cat: str, brand: str) -> str:
    if "MMO" in cat: return f"{brand}: CBN-licensed Mobile Money Operator (e-money/wallets)."
    if "PSSP" in cat: return f"{brand}: CBN-authorised Payment Solutions Service Provider (gateway/merchant collections)."
    if "PTSP" in cat: return f"{brand}: CBN-authorised Payment Terminal Service Provider (POS/terminals)."
    if "Super-Agent" in cat: return f"{brand}: CBN-authorised Super-Agent for agency banking distribution."
    if "Switching" in cat: return f"{brand}: CBN Switching and Processing licensee."
    if "IMTO" in cat: return f"{brand}: CBN-licensed International Money Transfer Operator."
    if "DML" in cat or "Money Lender" in cat: return f"{brand}: FCCPC-registered digital money lender / consumer credit apps."
    if "Card" in cat or "Scheme" in cat: return f"{brand}: CBN-listed card/payment scheme participant."
    if "Holding" in cat: return f"{brand}: CBN Payments Service Holding Company."
    if "PTSA" in cat: return f"{brand}: CBN Payments Terminal Service Aggregator."
    return f"{brand}: Nigeria tech/fintech company from a public regulator or directory source."

def parse_psp():
    cached = RAW / "parsed_psp.json"
    if cached.exists() and not (RAW / "psps.html").exists():
        return [(x["legal"], x["category"], x["source"]) for x in json.loads(cached.read_text())]
    html = (RAW / "psps.html").read_text(errors="ignore")
    sections = re.split(r"<h[1-4][^>]*>", html, flags=re.I)
    current = "CBN Payment Service Provider"
    out = []
    for block in sections:
        head = re.sub(r"<[^>]+>", " ", block[:160]).upper()
        if "CARD" in head and "SCHEME" in head: current = "CBN Card/Payment Scheme"
        elif "MOBILE MONEY" in head: current = "CBN Mobile Money Operator (MMO)"
        elif "SWITCHING" in head: current = "CBN Switching & Processing"
        elif "PSSP" in head or "PAYMENT SOLUTION SERVICE PROVIDER" in head: current = "CBN PSSP"
        elif "PTSP" in head or "PAYMENT TERMINAL SERVICES" in head: current = "CBN PTSP"
        elif "SUPER-AGENT" in head or "SUPER AGENT" in head: current = "CBN Super-Agent"
        elif "HOLDING" in head: current = "CBN Payments Service Holding"
        elif "AGGREGATOR" in head: current = "CBN PTSA"
        for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", block, flags=re.I|re.S):
            t = re.sub(r"<[^>]+>", "", c)
            t = re.sub(r"\s+", " ", t).strip().replace("&amp;", "&").replace("&nbsp;", " ")
            if not t or t.upper() in {"S/N","LICENCE","LICENCEE","SN"} or re.fullmatch(r"\d+\.?", t):
                continue
            if should_skip(t): continue
            out.append((t, current, "https://www.cbn.gov.ng/PaymentsSystem/PSPs.html"))
    return out

def parse_imto():
    cached = RAW / "parsed_imto.json"
    if cached.exists():
        return [(x["legal"], x["category"], x["source"]) for x in json.loads(cached.read_text())]
    out = []
    for n in (RAW / "imto_clean.txt").read_text().splitlines():
        if should_skip(n): continue
        if re.search(r"\b(Street|Road|Suite|Plaza|Building|Freeway|Lane|Close|Way)\b", n) and not re.search(r"\b(Limited|Ltd|Inc|PLC|LLC)\b", n, re.I):
            continue
        out.append((n, "CBN International Money Transfer Operator (IMTO)", "https://www.cbn.gov.ng/PaymentsSystem/InternationalMoneyTransferOperators.html"))
    return out

def parse_fccpc():
    cached = RAW / "parsed_fccpc.json"
    if cached.exists():
        return [(x["legal"], x["category"], x["source"]) for x in json.loads(cached.read_text())]
    out = []
    for n in (RAW / "fccpc_companies.txt").read_text().splitlines():
        n = n.strip().rstrip(",")
        if should_skip(n) or len(n) < 5: continue
        out.append((n, "FCCPC-registered Digital Money Lender (DML)", "https://fccpc.gov.ng/registration-of-digital-money-lenders/approvals-of-dmls/"))
    return out


def load_json_companies():
    out = []
    for f in sorted(DATA.glob("companies_part*.json")):
        out.extend(json.loads(f.read_text()))
    extra = DATA / "extra_funded.json"
    if extra.exists():
        out.extend(json.loads(extra.read_text()))
    return out

def main():
    funded = OrderedDict()
    def add_funded(item):
        key = norm_key(item["name"])
        if key in funded:
            if (item.get("funding") or 0) > (funded[key].get("funding") or 0):
                funded[key] = item
            return
        funded[key] = item
    for c in load_json_companies():
        add_funded(c)

    directory = OrderedDict()
    def add_dir(name, sector, product, hq, source, licences=""):
        if should_skip(name):
            return
        key = norm_key(name)
        if not key or key in funded:
            return
        if key in directory:
            d = directory[key]
            if product and len(product) > len(d["product"]):
                d["product"] = product
            if source and source not in d["sources"]:
                d["sources"] = (d["sources"] + " | " + source) if d["sources"] else source
            return
        ease = 2 if sector.lower() in {"fintech","insurtech","logistics","marketplace","healthtech","cleantech"} else 3
        if not licences:
            if sector.lower() in {"fintech","insurtech"}:
                licences = "CAC; NDPR; often CBN PSP/MMO/IMTO and/or FCCPC DML; SEC if capital-markets."
            elif sector.lower() == "healthtech":
                licences = "CAC; NDPR; MoH/state health rules; NAFDAC if drugs; payments via PSPs."
            else:
                licences = "CAC; NDPR."
        directory[key] = {
            "name": name, "sector": sector, "product": product, "hq": hq or "Nigeria",
            "founded": "", "ease": ease, "licences": licences, "sources": source,
            "notes": "No clean disclosed funding total in free sources used for this build; listed for directory coverage.",
        }

    curated = json.loads((DATA / "curated_directory.json").read_text())
    for c in curated:
        add_dir(c["name"], c["sector"], c["product"], c["hq"], c["source"])

    for legal, cat, src in parse_psp():
        brand = brand_from_legal(legal)
        add_dir(brand, "Fintech", product_for(cat, brand), "Nigeria", src, licences=f"CAC; NDPR; {cat}.")
    for legal, cat, src in parse_imto():
        brand = brand_from_legal(legal)
        add_dir(brand, "Fintech", product_for(cat, brand), "Nigeria", src, licences=f"CAC; NDPR; {cat}.")
    for legal, cat, src in parse_fccpc():
        brand = brand_from_legal(legal)
        add_dir(brand, "Fintech", product_for(cat, brand), "Nigeria", src, licences=f"CAC; NDPR; {cat}.")

    yc_path = RAW / "yc_nigeria_names.json"
    if yc_path.exists():
        for h in json.loads(yc_path.read_text()):
            add_dir(h["name"], "Other Tech", (h.get("desc") or "YC-backed company with Nigeria ops.")[:240], h.get("loc") or "Nigeria", "https://www.ycombinator.com/companies")
    sma = RAW / "sma_names.txt"
    if sma.exists():
        for n in sma.read_text().splitlines():
            if should_skip(n): continue
            add_dir(n, "Other Tech", f"{n}: Listed on StartupMapAfrica Nigeria sector directories.", "Nigeria", "https://startupmapafrica.com/startups/nigeria")

    funded_list = sorted(funded.values(), key=lambda x: x.get("funding") or 0, reverse=True)
    ranked_rows = []
    for i, c in enumerate(funded_list, 1):
        ranked_rows.append({
            "Rank": i, "Company": c["name"], "Sector": c["sector"], "Product_Summary": c["product"],
            "HQ_or_Primary_Market": c["hq"], "Founded_Year": c.get("founded") or "",
            "Total_Disclosed_Funding_USD": c.get("funding") if c.get("funding") is not None else "",
            "Valuation_USD": c["valuation"] if c.get("valuation") is not None else "Undisclosed",
            "Valuation_Status": c.get("val_status") or "Undisclosed",
            "Ease_to_Replicate": c.get("ease") or "", "Replicate_Notes": c.get("replicate") or "",
            "Licences_Regs_Needed": c.get("licences") or "", "Source_URLs": c.get("sources") or "",
            "Last_Verified_Date": TODAY,
        })

    dir_rows = []
    for d in directory.values():
        dir_rows.append({
            "Company": d["name"], "Sector": d["sector"], "Product_Summary": d["product"],
            "HQ_or_Primary_Market": d["hq"], "Founded_Year": d.get("founded") or "",
            "Total_Disclosed_Funding_USD": "", "Valuation_USD": "", "Valuation_Status": "Undisclosed",
            "Ease_to_Replicate": d.get("ease") or "",
            "Replicate_Notes": "Directory listing; funding undisclosed in free sources used.",
            "Licences_Regs_Needed": d.get("licences") or "", "Source_URLs": d.get("sources") or "",
            "Notes": d.get("notes") or "", "Last_Verified_Date": TODAY,
        })
    dir_rows.sort(key=lambda r: (r["Sector"], r["Company"].lower()))

    notable = [r for r in dir_rows if "CBN-" not in r["Product_Summary"] and "FCCPC-" not in r["Product_Summary"]]
    no_fund_rows = [{
        "Company": r["Company"], "Sector": r["Sector"], "Product_Summary": r["Product_Summary"],
        "HQ_or_Primary_Market": r["HQ_or_Primary_Market"], "Founded_Year": r["Founded_Year"],
        "Notes": r["Notes"], "Last_Verified_Date": TODAY,
    } for r in notable[:50]]

    ranked_fields = list(ranked_rows[0].keys())
    with open(DATA / "ranked_disclosed.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=ranked_fields); w.writeheader(); w.writerows(ranked_rows)
    dir_fields = list(dir_rows[0].keys())
    with open(DATA / "directory.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=dir_fields); w.writeheader(); w.writerows(dir_rows)
    with open(DATA / "no_disclosed_funding.csv", "w", newline="", encoding="utf-8") as f:
        fields = ["Company","Sector","Product_Summary","HQ_or_Primary_Market","Founded_Year","Notes","Last_Verified_Date"]
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(no_fund_rows)

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)

    def write_sheet(ws, rows, fields):
        for col, h in enumerate(fields, 1):
            cell = ws.cell(1, col, h); cell.fill = header_fill; cell.font = header_font
        for r_i, row in enumerate(rows, 2):
            for c_i, h in enumerate(fields, 1):
                ws.cell(r_i, c_i, row.get(h, ""))
        for col in range(1, len(fields)+1):
            ws.column_dimensions[get_column_letter(col)].width = 18

    ws1 = wb.active; ws1.title = "Ranked_Disclosed"
    write_sheet(ws1, ranked_rows, ranked_fields)
    ws2 = wb.create_sheet("Directory"); write_sheet(ws2, dir_rows, dir_fields)
    ws3 = wb.create_sheet("No_Disclosed_Funding")
    nf = ["Company","Sector","Product_Summary","HQ_or_Primary_Market","Founded_Year","Notes","Last_Verified_Date"]
    write_sheet(ws3, no_fund_rows, nf)
    ws4 = wb.create_sheet("Licences_Cheat_Sheet")
    with open(DATA / "licences_cheat_sheet.csv", newline="", encoding="utf-8") as f:
        reader = list(csv.reader(f))
    for r_i, row in enumerate(reader, 1):
        for c_i, val in enumerate(row, 1):
            cell = ws4.cell(r_i, c_i, val)
            if r_i == 1: cell.fill = header_fill; cell.font = header_font
    ws5 = wb.create_sheet("Methodology")
    method = [
        ["Field","Notes"],
        ["Ranking key","Total_Disclosed_Funding_USD descending; only rows with a sourced figure."],
        ["Directory","Companies without a clean disclosed funding total. Includes CBN PSP/MMO/IMTO and FCCPC DML licensees plus curated tech."],
        ["No fabrication","Funding and valuations never invented. Prefer blank / Undisclosed / Unknown."],
        ["Dedupe","Aliases merged (TeamApt=Moniepoint, Appzone=Zone, PayHippo=Rivy, Paylater=Carbon, Mkudi=Nomba, etc.)."],
        ["Scope","Nigeria-founded OR Nigeria primary market; fintech + tech platforms."],
        ["Ease_to_Replicate","1 hardest to 5 easiest for a small London digital consultancy builder profile."],
        ["Last verified", TODAY],
        ["Grand unique total", str(len(ranked_rows) + len(dir_rows))],
    ]
    write_sheet(ws5, [dict(zip(["Field","Notes"], r)) for r in method[1:]], ["Field","Notes"])
    ws5.cell(1,1,"Field").fill=header_fill; ws5.cell(1,1).font=header_font
    ws5.cell(1,2,"Notes").fill=header_fill; ws5.cell(1,2).font=header_font

    wb.save(DATA / "Nigeria_Fintech_Tech_Ranked.xlsx")
    wb.save(ROOT / "Nigeria_Fintech_Tech_Ranked.xlsx")

    meta = {
        "ranked_count": len(ranked_rows),
        "directory_count": len(dir_rows),
        "no_fund_notable_count": len(no_fund_rows),
        "grand_unique": len(ranked_rows) + len(dir_rows),
        "top10": [[r["Company"], r["Total_Disclosed_Funding_USD"]] for r in ranked_rows[:10]],
        "last_verified": TODAY,
    }
    (DATA / "build_meta.json").write_text(json.dumps(meta, indent=2))
    print(json.dumps(meta, indent=2))

if __name__ == "__main__":
    main()
